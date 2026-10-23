from django.contrib.auth import authenticate
from django.contrib.auth.models import User
from django.db.models import Sum, F, Q
from django.utils import timezone
from django.http import HttpResponse
from django.shortcuts import get_object_or_404

from rest_framework import viewsets, permissions, status
from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework.decorators import api_view, permission_classes, action
from rest_framework.authtoken.models import Token

import csv
import openpyxl
from io import BytesIO
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from .importers import read_rows, open_image_zip, import_rows, ImportFileError
from .models import UserProfile, Category, Size, Color, Product, ProductImage, Sale
from .serializers import (
    UserProfileSerializer, CategorySerializer, SizeSerializer, ColorSerializer,
    ProductImageSerializer, ProductReadSerializer, ProductWriteSerializer,
    SaleReadSerializer, SaleWriteSerializer,
)

import datetime
from decimal import Decimal
from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework.permissions import IsAuthenticated
from django.shortcuts import get_object_or_404
from .models import Cart, CartItem, OwnerNotification
from .serializers import CartItemSerializer, OwnerNotificationSerializer
from rest_framework.parsers import MultiPartParser, FormParser
from .importers import read_rows, import_rows, ImportFileError

# from .permissions import IsOwner   # you likely already import this

# =====================================================================
# Custom DRF Permissions
# =====================================================================

class IsOwnerOrReadOnly(permissions.BasePermission):
    def has_permission(self, request, view):
        if request.method in permissions.SAFE_METHODS:
            return True
        return (
            request.user and request.user.is_authenticated and
            hasattr(request.user, 'profile') and request.user.profile.role == 'OWNER'
        )


class IsOwner(permissions.BasePermission):
    def has_permission(self, request, view):
        return (
            request.user and request.user.is_authenticated and
            hasattr(request.user, 'profile') and request.user.profile.role == 'OWNER'
        )

# =====================================================================
# Auth APIs
# =====================================================================

@api_view(['POST'])
@permission_classes([permissions.AllowAny])
def api_login_view(request):
    username = request.data.get('username')
    password = request.data.get('password')
    user = authenticate(request, username=username, password=password)

    if user is None:
        return Response({'error': 'Invalid username or password.'}, status=401)

    token, _ = Token.objects.get_or_create(user=user)
    role = user.profile.role if hasattr(user, 'profile') else 'CUSTOMER'

    return Response({'token': token.key, 'username': user.username, 'role': role})


@api_view(['POST'])
@permission_classes([permissions.AllowAny])
def api_signup_view(request):
    username = request.data.get('username')
    email = request.data.get('email')
    password = request.data.get('password')
    role = request.data.get('role', 'CUSTOMER')

    if not username or not password or not email:
        return Response({'error': 'Please fill in all fields.'}, status=400)

    if User.objects.filter(username=username).exists():
        return Response({'error': 'Username already taken.'}, status=400)

    user = User.objects.create_user(username=username, email=email, password=password)
    user.profile.role = role if role in ('OWNER', 'CUSTOMER') else 'CUSTOMER'
    user.profile.save()

    token, _ = Token.objects.get_or_create(user=user)
    return Response({'token': token.key, 'username': user.username, 'role': user.profile.role}, status=201)


@api_view(['POST'])
@permission_classes([permissions.IsAuthenticated])
def api_logout_view(request):
    request.user.auth_token.delete()
    return Response({'message': 'Logged out successfully.'})


@api_view(['GET'])
@permission_classes([permissions.IsAuthenticated])
def api_current_user_view(request):
    return Response(UserProfileSerializer(request.user.profile).data)

# =====================================================================
# REST ViewSets
# =====================================================================

class ProductViewSet(viewsets.ModelViewSet):
    queryset = Product.objects.all().prefetch_related('sizes', 'colors', 'images', 'category')
    permission_classes = [IsOwnerOrReadOnly]

    def get_serializer_class(self):
        if self.action in ['list', 'retrieve']:
            return ProductReadSerializer
        return ProductWriteSerializer

    def get_queryset(self):
        queryset = self.queryset
        search = self.request.query_params.get('search')
        if search:
            queryset = queryset.filter(
                Q(name__icontains=search) | Q(description__icontains=search) | Q(fabric__icontains=search)
            )

        category_id = self.request.query_params.get('category')
        if category_id:
            try:
                cat = Category.objects.get(pk=category_id)
                subcats = self._get_subcategory_ids(cat)
                queryset = queryset.filter(category_id__in=subcats)
            except Category.DoesNotExist:
                pass

        size_id = self.request.query_params.get('size')
        if size_id:
            queryset = queryset.filter(sizes__id=size_id)

        color_id = self.request.query_params.get('color')
        if color_id:
            queryset = queryset.filter(colors__id=color_id)

        min_price = self.request.query_params.get('min_price')
        max_price = self.request.query_params.get('max_price')
        if min_price:
            queryset = queryset.filter(selling_price__gte=min_price)
        if max_price:
            queryset = queryset.filter(selling_price__lte=max_price)

        stock_status = self.request.query_params.get('stock_status')
        if stock_status == 'low':
            queryset = queryset.filter(stock__lte=F('low_stock_limit'), stock__gt=0)
        elif stock_status == 'out':
            queryset = queryset.filter(stock__lte=0)
        elif stock_status == 'in':
            queryset = queryset.filter(stock__gt=0)

        ordering = self.request.query_params.get('ordering', '-created_at')
        return queryset.order_by(ordering)

    def _get_subcategory_ids(self, category):
        ids = [category.id]
        for sub in category.subcategories.all():
            ids.extend(self._get_subcategory_ids(sub))
        return ids

    @action(detail=True, methods=['POST'], permission_classes=[IsOwner])
    def upload_images(self, request, pk=None):
        product = self.get_object()
        files = request.FILES.getlist('images')
        if not files:
            return Response({'error': 'No images provided'}, status=400)

        has_primary = product.images.filter(is_primary=True).exists()
        created_images = []
        for index, file in enumerate(files):
            is_pri = not has_primary and index == 0
            img = ProductImage.objects.create(product=product, image=file, is_primary=is_pri)
            created_images.append(ProductImageSerializer(img, context={'request': request}).data)

        return Response(created_images, status=201)

    @action(detail=True, methods=['DELETE'], url_path='delete_image/(?P<image_id>[0-9]+)', permission_classes=[IsOwner])
    def delete_image(self, request, pk=None, image_id=None):
        product = self.get_object()
        img = get_object_or_404(ProductImage, product=product, pk=image_id)
        img.image.delete()
        img.delete()
        return Response(status=204)

    @action(detail=False, methods=['POST'], url_path='import',permission_classes=[IsOwner], parser_classes=[MultiPartParser, FormParser])
    def import_products(self, request):
        upload = request.FILES.get('file')
        if not upload:
            return Response({'error': 'No file uploaded'}, status=400)
        if upload.size > 5 * 1024 * 1024:
            return Response({'error': 'File is larger than 5 MB'}, status=400)

        category_id = str(request.data.get('category', ''))
        category = Category.objects.filter(pk=category_id).first() if category_id.isdigit() else None
        if category is None:
            return Response({'error': 'Choose a section (category) for these products'}, status=400)

        try:
            rows = read_rows(upload)
            zip_upload = request.FILES.get('images_zip')
            image_pack = open_image_zip(zip_upload) if zip_upload else None
        except ImportFileError as exc:
            return Response({'error': str(exc)}, status=400)

        return Response(import_rows(rows, category, image_pack))


class CategoryViewSet(viewsets.ReadOnlyModelViewSet):
    queryset = Category.objects.filter(parent=None).prefetch_related('subcategories')
    serializer_class = CategorySerializer
    permission_classes = [permissions.AllowAny]


class SizeViewSet(viewsets.ReadOnlyModelViewSet):
    queryset = Size.objects.all()
    serializer_class = SizeSerializer
    permission_classes = [permissions.AllowAny]


class ColorViewSet(viewsets.ReadOnlyModelViewSet):
    queryset = Color.objects.all()
    serializer_class = ColorSerializer
    permission_classes = [permissions.AllowAny]


class SaleViewSet(viewsets.ModelViewSet):
    queryset = Sale.objects.all().order_by('-timestamp')
    permission_classes = [IsOwner]

    def get_serializer_class(self):
        if self.action in ['list', 'retrieve']:
            return SaleReadSerializer
        return SaleWriteSerializer

# =====================================================================
# Analytics
# =====================================================================

class AnalyticsView(APIView):
    permission_classes = [IsOwner]

    def get(self, request, *args, **kwargs):
        local_tz = timezone.get_current_timezone()
        now = timezone.now().astimezone(local_tz)
        today_start = now.replace(hour=0, minute=0, second=0, microsecond=0)
        month_start = now.replace(day=1, hour=0, minute=0, second=0, microsecond=0)

        total_products = Product.objects.count()
        total_inventory = Product.objects.aggregate(total=Sum('stock'))['total'] or 0
        low_stock_count = Product.objects.filter(stock__lte=F('low_stock_limit'), stock__gt=0).count()
        out_of_stock_count = Product.objects.filter(stock__lte=0).count()

        today_sales = Sale.objects.filter(timestamp__gte=today_start)
        today_revenue = today_sales.aggregate(total=Sum(F('quantity') * F('selling_price_snapshot')))['total'] or Decimal('0.00')
        today_profit = today_sales.aggregate(total=Sum('total_profit'))['total'] or Decimal('0.00')

        month_sales = Sale.objects.filter(timestamp__gte=month_start)
        month_revenue = month_sales.aggregate(total=Sum(F('quantity') * F('selling_price_snapshot')))['total'] or Decimal('0.00')
        month_profit = month_sales.aggregate(total=Sum('total_profit'))['total'] or Decimal('0.00')

        low_stock_items = Product.objects.filter(stock__lte=F('low_stock_limit')).order_by('stock')[:5]
        low_stock_data = ProductReadSerializer(low_stock_items, many=True, context={'request': request}).data

        thirty_days_ago = today_start - timezone.timedelta(days=30)
        sales_trend_data = Sale.objects.filter(timestamp__gte=thirty_days_ago).values('timestamp__date').annotate(
            revenue=Sum(F('quantity') * F('selling_price_snapshot')),
            profit=Sum('total_profit'),
            sales_count=Sum('quantity')
        ).order_by('timestamp__date')

        dates_dict = {s['timestamp__date']: s for s in sales_trend_data}
        sales_trend = []
        for d in range(30, -1, -1):
            date_check = (today_start - timezone.timedelta(days=d)).date()
            if date_check in dates_dict:
                sales_trend.append({
                    'date': date_check.strftime('%b %d'),
                    'revenue': float(dates_dict[date_check]['revenue']),
                    'profit': float(dates_dict[date_check]['profit']),
                    'count': dates_dict[date_check]['sales_count'],
                })
            else:
                sales_trend.append({'date': date_check.strftime('%b %d'), 'revenue': 0.0, 'profit': 0.0, 'count': 0})

        category_sales_data = Sale.objects.filter(timestamp__gte=month_start).values(
            'product__category__parent__name', 'product__category__name'
        ).annotate(
            revenue=Sum(F('quantity') * F('selling_price_snapshot')),
            profit=Sum('total_profit'),
            quantity_sold=Sum('quantity')
        ).order_by('-revenue')

        category_sales = []
        for c in category_sales_data:
            parent = c['product__category__parent__name']
            child = c['product__category__name']
            cat_name = f"{parent} - {child}" if parent else child
            category_sales.append({
                'category': cat_name,
                'revenue': float(c['revenue'] or 0),
                'profit': float(c['profit'] or 0),
                'quantity': c['quantity_sold'],
            })

        return Response({
            'total_products': total_products,
            'total_inventory': total_inventory,
            'low_stock_count': low_stock_count,
            'out_of_stock_count': out_of_stock_count,
            'today_revenue': float(today_revenue),
            'today_profit': float(today_profit),
            'month_revenue': float(month_revenue),
            'month_profit': float(month_profit),
            'low_stock_items': low_stock_data,
            'sales_trend': sales_trend,
            'category_sales': category_sales,
        })

# =====================================================================
# Report Exporter
# =====================================================================

class ExportReportView(APIView):
    permission_classes = [IsOwner]

    def get(self, request, *args, **kwargs):
        report_type = request.query_params.get('report_type', 'monthly')
        start_date_str = request.query_params.get('start_date')
        end_date_str = request.query_params.get('end_date')
        export_format = request.query_params.get('export_format', 'csv')

        local_tz = timezone.get_current_timezone()
        now = timezone.now().astimezone(local_tz)
        today_start = now.replace(hour=0, minute=0, second=0, microsecond=0)

        if report_type == 'daily':
            start_date = today_start
            end_date = now
            title = f"Daily Sales Report - {start_date.strftime('%Y-%m-%d')}"
        elif report_type == 'weekly':
            start_date = today_start - timezone.timedelta(days=7)
            end_date = now
            title = f"Weekly Sales Report ({start_date.strftime('%Y-%m-%d')} to {end_date.strftime('%Y-%m-%d')})"
        elif report_type == 'monthly':
            start_date = today_start.replace(day=1)
            end_date = now
            title = f"Monthly Sales Report - {start_date.strftime('%B %Y')}"
        elif report_type == 'yearly':
            start_date = today_start.replace(month=1, day=1)
            end_date = now
            title = f"Yearly Sales Report - {start_date.strftime('%Y')}"
        elif report_type == 'custom' and start_date_str and end_date_str:
            try:
                start_date = timezone.make_aware(datetime.datetime.strptime(start_date_str, '%Y-%m-%d'), local_tz)
                end_date = timezone.make_aware(datetime.datetime.strptime(end_date_str, '%Y-%m-%d'), local_tz) + timezone.timedelta(days=1) - timezone.timedelta(seconds=1)
                title = f"Custom Sales Report ({start_date.strftime('%Y-%m-%d')} to {end_date_str})"
            except ValueError:
                return Response({'error': 'Invalid date format. Use YYYY-MM-DD'}, status=400)
        else:
            start_date = today_start.replace(day=1)
            end_date = now
            title = f"Monthly Sales Report - {start_date.strftime('%B %Y')}"

        sales = Sale.objects.filter(timestamp__range=(start_date, end_date)).order_by('timestamp')

        total_revenue = Decimal('0.00')
        total_profit = Decimal('0.00')
        total_qty = 0
        for sale in sales:
            total_revenue += sale.selling_price_snapshot * sale.quantity
            total_profit += sale.total_profit
            total_qty += sale.quantity

        best_selling = sales.values('product_name_snapshot').annotate(
            qty=Sum('quantity'),
            profit=Sum('total_profit'),
            revenue=Sum(F('quantity') * F('selling_price_snapshot'))
        ).order_by('-qty')[:5]

        category_summary = sales.values('product__category__parent__name', 'product__category__name').annotate(
            qty=Sum('quantity'),
            revenue=Sum(F('quantity') * F('selling_price_snapshot')),
            profit=Sum('total_profit')
        ).order_by('-revenue')

        if export_format == 'json':
            best_selling_list = [{
                'product_name_snapshot': b['product_name_snapshot'], 'qty': b['qty'],
                'revenue': float(b['revenue'] or 0), 'profit': float(b['profit'] or 0)
            } for b in best_selling]

            category_summary_list = []
            for c in category_summary:
                parent = c['product__category__parent__name']
                child = c['product__category__name']
                cat_name = f"{parent} -> {child}" if parent else child
                category_summary_list.append({
                    'category': cat_name, 'qty': c['qty'],
                    'revenue': float(c['revenue'] or 0), 'profit': float(c['profit'] or 0)
                })

            return Response({
                'title': title, 'total_revenue': float(total_revenue), 'total_profit': float(total_profit),
                'total_qty': total_qty, 'best_selling': best_selling_list, 'category_summary': category_summary_list
            })
        elif export_format == 'csv':
            return self._export_csv(sales, title, total_revenue, total_profit, total_qty, best_selling)
        elif export_format == 'excel':
            return self._export_excel(sales, title, total_revenue, total_profit, total_qty, best_selling, category_summary)
        elif export_format == 'pdf':
            return self._export_pdf(sales, title, total_revenue, total_profit, total_qty, best_selling, category_summary, start_date, end_date)
        else:
            return Response({'error': 'Unsupported export format'}, status=400)

    def _export_csv(self, sales, title, total_rev, total_prof, total_qty, best_selling):
        response = HttpResponse(content_type='text/csv')
        response['Content-Disposition'] = f'attachment; filename="report_{timezone.now().strftime("%Y%m%d_%H%M%S")}.csv"'
        writer = csv.writer(response)
        writer.writerow([title])
        writer.writerow([])
        writer.writerow(['Report Summary'])
        writer.writerow(['Total Revenue', f"{total_rev:.2f}"])
        writer.writerow(['Total Profit', f"{total_prof:.2f}"])
        writer.writerow(['Total Quantity Sold', total_qty])
        writer.writerow([])
        writer.writerow(['Transactions Log'])
        writer.writerow(['Date & Time', 'Product Name', 'Quantity', 'Selling Price', 'Base Price', 'Total Profit', 'Payment Method', 'Customer Name'])
        for s in sales:
            writer.writerow([
                s.timestamp.strftime('%Y-%m-%d %H:%M'), s.product_name_snapshot, s.quantity,
                f"{s.selling_price_snapshot:.2f}", f"{s.base_price_snapshot:.2f}",
                f"{s.total_profit:.2f}", s.payment_method, s.customer_name or 'N/A'
            ])
        writer.writerow([])
        writer.writerow(['Top Selling Products'])
        writer.writerow(['Product Name', 'Quantity Sold', 'Revenue Generated', 'Profit Earned'])
        for b in best_selling:
            writer.writerow([b['product_name_snapshot'], b['qty'], f"{b['revenue']:.2f}", f"{b['profit']:.2f}"])
        return response

    def _export_excel(self, sales, title, total_rev, total_prof, total_qty, best_selling, category_summary):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sales Transactions"
        ws.merge_cells('A1:H1')
        title_cell = ws['A1']
        title_cell.value = title
        title_cell.font = openpyxl.styles.Font(size=16, bold=True, color="1F4E79")
        title_cell.alignment = openpyxl.styles.Alignment(horizontal="center")

        ws['A3'] = "Financial Summary"
        ws['A3'].font = openpyxl.styles.Font(bold=True, size=12)
        ws['A4'] = "Total Revenue"
        ws['B4'] = float(total_rev)
        ws['A5'] = "Total Profit"
        ws['B5'] = float(total_prof)
        ws['A6'] = "Total Products Sold"
        ws['B6'] = total_qty
        ws['B4'].number_format = '$#,##0.00'
        ws['B5'].number_format = '$#,##0.00'

        headers = ['Date & Time', 'Product Name', 'Quantity', 'Selling Price', 'Base Price', 'Total Profit', 'Payment Method', 'Customer Name']
        ws.append([])
        ws.append([])
        ws.append(headers)
        header_font = openpyxl.styles.Font(bold=True, color="FFFFFF")
        header_fill = openpyxl.styles.PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        for col_num in range(1, 9):
            cell = ws.cell(row=9, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = openpyxl.styles.Alignment(horizontal="center")

        current_row = 10
        for s in sales:
            ws.append([
                s.timestamp.strftime('%Y-%m-%d %H:%M'), s.product_name_snapshot, s.quantity,
                float(s.selling_price_snapshot), float(s.base_price_snapshot), float(s.total_profit),
                s.payment_method, s.customer_name or 'N/A'
            ])
            ws.cell(row=current_row, column=4).number_format = '$#,##0.00'
            ws.cell(row=current_row, column=5).number_format = '$#,##0.00'
            ws.cell(row=current_row, column=6).number_format = '$#,##0.00'
            current_row += 1

        ws2 = wb.create_sheet(title="Analytics & Categories")
        ws2['A1'] = "Top Performing Products"
        ws2['A1'].font = openpyxl.styles.Font(bold=True, size=13)
        ws2.append(['Product Name', 'Quantity Sold', 'Revenue', 'Profit'])
        for col in range(1, 5):
            ws2.cell(row=2, column=col).font = openpyxl.styles.Font(bold=True)

        row_idx = 3
        for b in best_selling:
            ws2.append([b['product_name_snapshot'], b['qty'], float(b['revenue']), float(b['profit'])])
            ws2.cell(row=row_idx, column=3).number_format = '$#,##0.00'
            ws2.cell(row=row_idx, column=4).number_format = '$#,##0.00'
            row_idx += 1

        ws2.append([])
        ws2.append([])
        cat_start_row = row_idx + 2
        ws2.cell(row=cat_start_row, column=1, value="Category-wise Performance").font = openpyxl.styles.Font(bold=True, size=13)
        ws2.cell(row=cat_start_row+1, column=1, value="Category").font = openpyxl.styles.Font(bold=True)
        ws2.cell(row=cat_start_row+1, column=2, value="Quantity Sold").font = openpyxl.styles.Font(bold=True)
        ws2.cell(row=cat_start_row+1, column=3, value="Revenue").font = openpyxl.styles.Font(bold=True)
        ws2.cell(row=cat_start_row+1, column=4, value="Profit").font = openpyxl.styles.Font(bold=True)

        current_cat_row = cat_start_row + 2
        for c in category_summary:
            parent = c['product__category__parent__name']
            child = c['product__category__name']
            cat_name = f"{parent} - {child}" if parent else child
            ws2.append([cat_name, c['qty'], float(c['revenue']), float(c['profit'])])
            ws2.cell(row=current_cat_row, column=3).number_format = '$#,##0.00'
            ws2.cell(row=current_cat_row, column=4).number_format = '$#,##0.00'
            current_cat_row += 1

        for sheet in wb.worksheets:
            for col in sheet.columns:
                max_len = 0
                col_letter = openpyxl.utils.get_column_letter(col[0].column)
                for cell in col:
                    if cell.value:
                        max_len = max(max_len, len(str(cell.value)))
                sheet.column_dimensions[col_letter].width = max(max_len + 3, 10)

        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = f'attachment; filename="report_{timezone.now().strftime("%Y%m%d_%H%M%S")}.xlsx"'
        wb.save(response)
        return response

    def _export_pdf(self, sales, title, total_rev, total_prof, total_qty, best_selling, category_summary, start_date, end_date):
        buffer = BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
        story = []
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(name='TitleStyle', fontName='Helvetica-Bold', fontSize=22,
                                      textColor=colors.HexColor('#1F4E79'), spaceAfter=15, alignment=1)
        subtitle_style = ParagraphStyle(name='SubtitleStyle', fontName='Helvetica', fontSize=11,
                                         textColor=colors.HexColor('#555555'), spaceAfter=20, alignment=1)
        heading_style = ParagraphStyle(name='HeadingStyle', fontName='Helvetica-Bold', fontSize=14,
                                        textColor=colors.HexColor('#2E4053'), spaceBefore=15, spaceAfter=10)
        body_style = ParagraphStyle(name='BodyStyle', fontName='Helvetica', fontSize=9, spaceAfter=5)

        story.append(Paragraph("Cloth Shop Management System", title_style))
        story.append(Paragraph(f"{title}<br/>Generated on: {timezone.now().strftime('%Y-%m-%d %I:%M %p')}", subtitle_style))
        story.append(Spacer(1, 10))

        summary_data = [
            [Paragraph("<b>Financial Summary Indicator</b>", body_style), Paragraph("<b>Value</b>", body_style)],
            ["Total Revenue", f"INR {total_rev:,.2f}"],
            ["Total Profit", f"INR {total_prof:,.2f}"],
            ["Total Products Sold", str(total_qty)],
            ["Net Profit Margin", f"{(total_prof / total_rev * 100 if total_rev > 0 else 0):.1f}%"]
        ]
        t_summary = Table(summary_data, colWidths=[3.0*inch, 2.5*inch])
        t_summary.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#EAECEE')),
            ('ALIGN', (0,0), (-1,-1), 'LEFT'),
            ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
            ('GRID', (0,0), (-1,-1), 1, colors.HexColor('#BDC3C7')),
            ('PADDING', (0,0), (-1,-1), 6),
        ]))
        story.append(Paragraph("Financial Performance Overview", heading_style))
        story.append(t_summary)
        story.append(Spacer(1, 20))

        story.append(Paragraph("Top 5 Best-Selling Products", heading_style))
        best_data = [[Paragraph("<b>Product Name</b>", body_style), Paragraph("<b>Qty Sold</b>", body_style),
                      Paragraph("<b>Revenue</b>", body_style), Paragraph("<b>Profit</b>", body_style)]]
        for b in best_selling:
            best_data.append([b['product_name_snapshot'], str(b['qty']), f"INR {b['revenue']:,.2f}", f"INR {b['profit']:,.2f}"])
        t_best = Table(best_data, colWidths=[3.2*inch, 1.0*inch, 1.6*inch, 1.6*inch])
        t_best.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#E5E8E8')),
            ('ALIGN', (0,0), (0,-1), 'LEFT'), ('ALIGN', (1,0), (-1,-1), 'RIGHT'),
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#BDC3C7')), ('PADDING', (0,0), (-1,-1), 5),
        ]))
        story.append(t_best)
        story.append(Spacer(1, 20))

        story.append(Paragraph("Category Performance Summary", heading_style))
        cat_data = [[Paragraph("<b>Category</b>", body_style), Paragraph("<b>Qty Sold</b>", body_style),
                     Paragraph("<b>Revenue</b>", body_style), Paragraph("<b>Profit</b>", body_style)]]
        for c in category_summary:
            parent = c['product__category__parent__name']
            child = c['product__category__name']
            cat_name = f"{parent} -> {child}" if parent else child
            cat_data.append([cat_name, str(c['qty']), f"INR {c['revenue']:,.2f}", f"INR {c['profit']:,.2f}"])
        t_cat = Table(cat_data, colWidths=[3.2*inch, 1.0*inch, 1.6*inch, 1.6*inch])
        t_cat.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#E5E8E8')),
            ('ALIGN', (0,0), (0,-1), 'LEFT'), ('ALIGN', (1,0), (-1,-1), 'RIGHT'),
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#BDC3C7')), ('PADDING', (0,0), (-1,-1), 5),
        ]))
        story.append(t_cat)
        story.append(PageBreak())

        story.append(Paragraph("Detailed Transactions Log", heading_style))
        log_data = [[Paragraph("<b>Date/Time</b>", body_style), Paragraph("<b>Product Name</b>", body_style),
                     Paragraph("<b>Qty</b>", body_style), Paragraph("<b>Price</b>", body_style),
                     Paragraph("<b>Profit</b>", body_style), Paragraph("<b>Method</b>", body_style)]]
        for s in sales:
            log_data.append([
                s.timestamp.strftime('%y-%m-%d %H:%M'),
                s.product_name_snapshot[:28] + '...' if len(s.product_name_snapshot) > 30 else s.product_name_snapshot,
                str(s.quantity), f"{s.selling_price_snapshot:,.0f}", f"{s.total_profit:,.0f}", s.payment_method
            ])
        t_log = Table(log_data, colWidths=[1.1*inch, 3.2*inch, 0.5*inch, 0.9*inch, 0.9*inch, 0.8*inch])
        t_log.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#EAECEE')),
            ('ALIGN', (0,0), (1,-1), 'LEFT'), ('ALIGN', (2,0), (4,-1), 'RIGHT'), ('ALIGN', (5,0), (5,-1), 'CENTER'),
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#D5D8DC')),
            ('FONTSIZE', (0,0), (-1,-1), 8), ('PADDING', (0,0), (-1,-1), 4),
        ]))
        story.append(t_log)

        doc.build(story)
        pdf_content = buffer.getvalue()
        buffer.close()

        response = HttpResponse(content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="report_{timezone.now().strftime("%Y%m%d_%H%M%S")}.pdf"'
        response.write(pdf_content)
        return response

class CartView(APIView):
    """GET /api/cart/ : the logged-in customer's cart"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        cart, _ = Cart.objects.get_or_create(user=request.user)
        items = cart.items.select_related('product')
        return Response(CartItemSerializer(items, many=True).data)


class AddToCartView(APIView):
    """POST /api/cart/add/ with {product_id, quantity}"""
    permission_classes = [IsAuthenticated]

    def post(self, request):
        # 1. Read and validate input
        try:
            quantity = int(request.data.get('quantity', 1))
        except (TypeError, ValueError):
            return Response({'error': 'Quantity must be a number'}, status=400)
        if quantity < 1:
            return Response({'error': 'Quantity must be at least 1'}, status=400)

        product = get_object_or_404(Product, pk=request.data.get('product_id'))

        # 2. Stock checks
        if product.is_out_of_stock:
            return Response({'error': 'This product is out of stock'}, status=400)

        # 3. Find or create this customer's cart and line item
        cart, _ = Cart.objects.get_or_create(user=request.user)
        item, _ = CartItem.objects.get_or_create(cart=cart, product=product, defaults={'quantity': 0})

        new_quantity = item.quantity + quantity
        if new_quantity > product.stock:
            return Response({'error': f'Only {product.stock} units available'}, status=400)

        item.quantity = new_quantity
        item.save()

        # 4. Tell the owner
        OwnerNotification.objects.create(
            customer=request.user,
            product=product,
            message=f"{request.user.username} added {quantity} x {product.name} to cart",
        )

        return Response(CartItemSerializer(item).data, status=201)


class CartItemDeleteView(APIView):
    """DELETE /api/cart/items/<id>/ : remove a line from my cart"""
    permission_classes = [IsAuthenticated]

    def delete(self, request, item_id):
        item = get_object_or_404(CartItem, pk=item_id, cart__user=request.user)
        item.delete()
        return Response(status=204)


class OwnerNotificationListView(APIView):
    """GET /api/owner/notifications/ : latest notifications + unread count"""
    permission_classes = [IsOwner]

    def get(self, request):
        latest = OwnerNotification.objects.all()[:20]
        unread = OwnerNotification.objects.filter(is_read=False).count()
        return Response({
            'unread_count': unread,
            'results': OwnerNotificationSerializer(latest, many=True).data,
        })


class OwnerNotificationMarkReadView(APIView):
    """POST /api/owner/notifications/mark-read/"""
    permission_classes = [IsOwner]

    def post(self, request):
        OwnerNotification.objects.filter(is_read=False).update(is_read=True)
        return Response({'status': 'ok'})

    