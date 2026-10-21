import csv
import io
from decimal import Decimal, InvalidOperation

from django.db import transaction
from openpyxl import load_workbook

from .models import Product, Size, Color

MAX_ROWS = 1000


class ImportFileError(Exception):
    """The whole file is unusable (wrong type, missing columns...)."""


def _clean(value):
    return '' if value is None else str(value).strip()


def read_rows(upload):
    """Return [(line_number, {column: value}), ...] from a .csv or .xlsx upload."""
    filename = upload.name.lower()

    if filename.endswith('.csv'):
        try:
            text = upload.read().decode('utf-8-sig')      # utf-8-sig removes Excel's hidden BOM
        except UnicodeDecodeError:
            raise ImportFileError('Save the CSV as "CSV UTF-8" and try again')
        raw = list(csv.reader(io.StringIO(text)))
    elif filename.endswith('.xlsx'):
        try:
            workbook = load_workbook(upload, read_only=True, data_only=True)
            raw = [list(r) for r in workbook.active.iter_rows(values_only=True)]
        except Exception:
            raise ImportFileError('Could not read this Excel file')
    else:
        raise ImportFileError('Upload a .csv or .xlsx file')

    if not raw:
        raise ImportFileError('The file is empty')

    headers = [_clean(h).lower().replace(' ', '_') for h in raw[0]]
    missing = [h for h in ('name', 'selling_price') if h not in headers]
    if missing:
        raise ImportFileError(f"Missing column(s): {', '.join(missing)}")

    rows = []
    for line, values in enumerate(raw[1:], start=2):       # line 2 = first data row
        record = {h: _clean(v) for h, v in zip(headers, values)}
        if any(record.values()):                           # skip blank lines
            rows.append((line, record))

    if len(rows) > MAX_ROWS:
        raise ImportFileError(f'Too many rows (limit is {MAX_ROWS} per file)')
    return rows


def _decimal(value, field, required=False):
    if value == '':
        if required:
            raise ValueError(f'{field} is required')
        return Decimal('0')
    try:
        number = Decimal(value)
    except InvalidOperation:
        raise ValueError(f'{field} must be a number (got "{value}")')
    if number < 0:
        raise ValueError(f'{field} cannot be negative')
    return number


def _int(value, field, default):
    if value == '':
        return default
    try:
        number = int(float(value))          # Excel often gives 20.0
    except ValueError:
        raise ValueError(f'{field} must be a whole number (got "{value}")')
    if number < 0:
        raise ValueError(f'{field} cannot be negative')
    return number


def _split(value):
    return [part.strip() for part in value.split(',') if part.strip()]


def _parse(record):
    """Validate one row. Returns (product_fields, size_names, color_specs)."""
    name = record.get('name', '')
    if not name:
        raise ValueError('name is required')
    if len(name) > 200:
        raise ValueError('name is longer than 200 characters')

    fields = {
        'name': name,
        'description': record.get('description', ''),
        'selling_price': _decimal(record.get('selling_price', ''), 'selling_price', required=True),
        'base_price': _decimal(record.get('base_price', ''), 'base_price'),
        'stock': _int(record.get('stock', ''), 'stock', 0),
        'low_stock_limit': _int(record.get('low_stock_limit', ''), 'low_stock_limit', 5),
        'fabric': record.get('fabric', '')[:100],
    }
    return fields, _split(record.get('sizes', '')), _split(record.get('colors', ''))


def import_rows(rows, category):
    created = updated = 0
    errors = []

    for line, record in rows:
        try:
            fields, size_names, color_specs = _parse(record)

            with transaction.atomic():          # one row = all-or-nothing
                product = Product.objects.filter(name__iexact=fields['name'], category=category).first()
                is_new = product is None

                if is_new:
                    product = Product.objects.create(category=category, **fields)
                else:
                    for key, value in fields.items():
                        setattr(product, key, value)
                    product.save()

                # Only touch sizes/colors if the row actually has them
                if size_names:
                    product.sizes.set([Size.objects.get_or_create(name=n)[0] for n in size_names])
                if color_specs:
                    colors = []
                    for spec in color_specs:
                        color_name, _, hex_code = spec.partition(':')
                        color, _ = Color.objects.get_or_create(
                            name=color_name.strip(),
                            defaults={'hex_code': hex_code.strip() or '#cccccc'},
                        )
                        colors.append(color)
                    product.colors.set(colors)

            created += is_new
            updated += not is_new

        except ValueError as exc:
            errors.append({'row': line, 'error': str(exc)})

    return {'created': created, 'updated': updated, 'skipped': len(errors), 'errors': errors}
